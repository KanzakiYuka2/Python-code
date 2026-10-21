import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.drawing.image import Image


# =========================================================
# 1. Settings
# =========================================================
# 入力CSVファイル名
input_path = Path("Test202606220001.CSV")

# 出力Excelファイル名
output_excel_path = Path("Test202606220001_average.xlsx")

# 出力するグラフ画像名
h_graph_path = Path("leakrate_timedot0622avg.png")
i_graph_path = Path("pressure_timedot0622avg.png")

# 平均化する秒数
average_sec = 10


# =========================================================
# 2. Excel-like ROUND function
# =========================================================
# ExcelのROUND関数に近い丸め方
# Python標準のround()はExcelと少し違うため、自作する
def excel_round(x, digits=0):
    factor = 10 ** digits
    return np.sign(x) * np.floor(np.abs(x) * factor + 0.5) / factor


# =========================================================
# 3. Read CSV
# =========================================================
# このCSVは1〜14行目がメタ情報
# 15行目から実際のデータなので skiprows=14 にする
# encoding は日本語CSV対策で shift_jis を指定

df = pd.read_csv(
    input_path,
    skiprows=14,
    header=None,
    names=["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"],
    encoding="shift_jis"
)


# =========================================================
# 4. Convert columns to numeric
# =========================================================
# 文字列として読まれた数字を、計算できる数値型に変換する

for col in ["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"]:
    df[col] = pd.to_numeric(df[col], errors="coerce")

# 読み取りに失敗した行を削除
df = df.dropna(
    subset=["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"]
).reset_index(drop=True)


# =========================================================
# 5. Calculate F to K
# =========================================================
# Excelの F16:K16 で行っていた計算をPythonで再現する

# F = 10 - C - 12
df["F"] = 10 - df["C_CH2"] - 12

# G = ROUND(F, 0)
df["G"] = excel_round(df["F"], 0)

# H = B × 10^G
# H列：リークレート相当
df["H"] = df["B_CH1"] * (10 ** df["G"])

# I = 10^(D - 5.5) × 100
# I列：圧力相当
df["I"] = (10 ** (df["D_CH3"] - 5.5)) * 100

# J = 10^(E - 3)
df["J"] = 10 ** (df["E_CH4"] - 3)

# K = J - I
df["K"] = df["J"] - df["I"]


# =========================================================
# 6. Create 10-second average data
# =========================================================
# time_sを10秒ごとに区切って平均を取る
# 例：0〜9秒、10〜19秒、20〜29秒...

df_10s = (
    df
    .assign(time_bin=(df["time_s"] // average_sec) * average_sec)
    .groupby("time_bin", as_index=False)
    .agg({
        "B_CH1": "mean",
        "C_CH2": "mean",
        "D_CH3": "mean",
        "E_CH4": "mean",
        "F": "mean",
        "G": "mean",
        "H": "mean",
        "I": "mean",
        "J": "mean",
        "K": "mean"
    })
)

print("10-second average data created.")


# =========================================================
# 7. Save processed data to Excel
# =========================================================
# Excelに2つのシートを作る
# 1. Processed Data：元データ＋F〜K計算結果
# 2. 10s Average：10秒平均データ

with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Processed Data")
    df_10s.to_excel(writer, index=False, sheet_name="10s Average")

print(f"Processed Excel saved: {output_excel_path}")


# =========================================================
# 8. Dot graph function
# =========================================================
# ドットのみのグラフを作る関数
# 以前のラインプロット plt.plot() ではなく、
# plt.scatter() を使うことで点だけのグラフにする

def make_dot_graph(data, x_col, y_col, title, ylabel, save_path):
    # 必要な列だけ取り出す
    plot_data = data[[x_col, y_col]].copy()

    # NaNを削除
    plot_data = plot_data.dropna()

    # logスケールでは0以下の値を表示できないので削除
    plot_data = plot_data[plot_data[y_col] > 0]

    plt.figure(figsize=(10, 5))

    # ドットのみのグラフ
    plt.scatter(
        plot_data[x_col],
        plot_data[y_col],
        s=8
    )

    plt.xlabel("Time [s]")
    plt.ylabel(ylabel)
    plt.title(title)
    plt.grid(True)

    # リークレートや圧力は桁が大きく変わるのでlog表示
    plt.yscale("log")

    # 文字や軸ラベルが切れないように調整
    plt.tight_layout()

    # PNGとして保存
    plt.savefig(save_path, dpi=300)
    plt.close()

    print(f"Graph saved: {save_path}")


# =========================================================
# 9. H column 10-second average dot graph
# =========================================================
# H列の10秒平均をドットグラフにする

make_dot_graph(
    data=df_10s,
    x_col="time_bin",
    y_col="H",
    title="Leak Rate vs Time 10-s Average",
    ylabel="Leak rate",
    save_path=h_graph_path
)


# =========================================================
# 10. I column 10-second average dot graph
# =========================================================
# I列の10秒平均をドットグラフにする

make_dot_graph(
    data=df_10s,
    x_col="time_bin",
    y_col="I",
    title="Pressure vs Time 10-s Average",
    ylabel="Pressure",
    save_path=i_graph_path
)


# =========================================================
# 11. Insert graphs into Excel
# =========================================================
# 保存したPNGグラフをExcelのGraphsシートに貼り付ける

wb = load_workbook(output_excel_path)

# Graphsシートを作成
if "Graphs" in wb.sheetnames:
    ws_graph = wb["Graphs"]
else:
    ws_graph = wb.create_sheet("Graphs")

# グラフタイトルをExcel内に書く
ws_graph["A1"] = "Leak Rate vs Time 10-s Average"
ws_graph["A25"] = "Pressure vs Time 10-s Average"

# PNG画像を読み込む
img_h = Image(str(h_graph_path))
img_i = Image(str(i_graph_path))

# Excel内での画像サイズを調整
img_h.width = 720
img_h.height = 360

img_i.width = 720
img_i.height = 360

# 画像を貼り付け
ws_graph.add_image(img_h, "A2")
ws_graph.add_image(img_i, "A26")

# Excelを保存
wb.save(output_excel_path)

print(f"Graphs inserted into Excel: {output_excel_path}")