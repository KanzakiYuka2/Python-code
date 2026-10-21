import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.drawing.image import Image


# =========================================================
# Settings
# =========================================================
    
input_path = Path("Test202606170001.CSV")
output_excel_path = Path("Test202606170001_0p2Hz_30point_average.xlsx")

h_graph_path = Path("leakrate_time0617_0p2Hz_30avg.png")
i_graph_path = Path("pressure_time0617_0p2Hz_30avg.png")

# 元データは1 Hz
# 5個ごとに1個取る → 0.2 Hz
sampling_step = 5

# 0.2 Hz化した後、30個ごとに平均
average_points = 30


# =========================================================
# Excel-like ROUND function
# =========================================================

def excel_round(x, digits=0):
    """
    ExcelのROUND関数に近い丸め方。
    Python標準のround()とは丸め方が違う場合があるため、自作する。
    """
    factor = 10 ** digits
    return np.sign(x) * np.floor(np.abs(x) * factor + 0.5) / factor


# =========================================================
# Read CSV
# =========================================================
# いつものCSV：
# 1〜14行目がメタ情報
# 15行目から実データ

df = pd.read_csv(
    input_path,
    skiprows=14,
    header=None,
    names=["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"],
    encoding="shift_jis"
)


# =========================================================
# Convert columns to numeric
# =========================================================

for col in ["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"]:
    df[col] = pd.to_numeric(df[col], errors="coerce")

df = df.dropna(
    subset=["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"]
).reset_index(drop=True)


# =========================================================
# F to K calculation
# Same logic as Excel F16:K16
# =========================================================

df["F"] = 10 - df["C_CH2"] - 12
df["G"] = excel_round(df["F"], 0)
df["H"] = df["B_CH1"] * (10 ** df["G"])
df["I"] = (10 ** (df["D_CH3"] - 5.5)) * 100
df["J"] = 10 ** (df["E_CH4"] - 3)
df["K"] = df["J"] - df["I"]


# =========================================================
# Downsample to 0.2 Hz
# =========================================================
# 元データが1 Hzなので、5個ごとに1個選ぶ。
# pandasでは0始まりなので、
# iloc[::5] = 1番目, 6番目, 11番目, 16番目...
# に相当する。

df_02hz = df.iloc[::sampling_step].copy().reset_index(drop=True)

# 確認用の列を追加
df_02hz["selected_point_no"] = df_02hz.index + 1

print(f"Original data points: {len(df)}")
print(f"0.2 Hz selected data points: {len(df_02hz)}")


# =========================================================
# Average every 30 selected points
# =========================================================
# 0.2 Hzのデータを30個ごとに平均する。
# 0.2 Hz = 5秒ごとに1点
# 30点平均 = 150秒ごとの平均

df_02hz["avg_group"] = df_02hz.index // average_points

df_avg = (
    df_02hz
    .groupby("avg_group", as_index=False)
    .agg({
        "time_s": "mean",
        "B_CH1": "mean",
        "C_CH2": "mean",
        "D_CH3": "mean",
        "E_CH4": "mean",
        "F": "mean",
        "G": "mean",
        "H": "mean",
        "I": "mean",
        "J": "mean",
        "K": "mean",
        "selected_point_no": "count"
    })
)

# 列名を分かりやすく変更
df_avg = df_avg.rename(columns={
    "selected_point_no": "points_in_average",
    "time_s": "time_s_mean"
})

# 時間をhourでも追加
df["time_h"] = df["time_s"] / 3600
df_02hz["time_h"] = df_02hz["time_s"] / 3600
df_avg["time_h_mean"] = df_avg["time_s_mean"] / 3600

print(f"Average data points: {len(df_avg)}")
print(f"Each average point is about {sampling_step * average_points} seconds.")


# =========================================================
# Save processed data to Excel
# =========================================================

with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Processed Data 1Hz")
    df_02hz.to_excel(writer, index=False, sheet_name="Selected 0.2Hz")
    df_avg.to_excel(writer, index=False, sheet_name="30point Average")

print(f"Processed Excel saved: {output_excel_path}")


# =========================================================
# Graph function
# =========================================================

def make_average_graph(data, x_col, y_col, title, ylabel, save_path):
    """
    平均値グラフを作る関数。
    今回は平均値同士の変化を見やすくするため、line + marker にする。
    ドットだけにしたい場合は plt.plot を plt.scatter に変更する。
    """

    plot_data = data[[x_col, y_col]].copy()
    plot_data = plot_data.dropna()

    # log scaleでは0以下を表示できないので除外
    plot_data = plot_data[plot_data[y_col] > 0]

    plt.figure(figsize=(10, 5))

    # 平均値グラフ：線 + 点
    plt.plot(
        plot_data[x_col],
        plot_data[y_col],
        marker="o",
        markersize=3,
        linewidth=1
    )

    plt.xlabel("Time [s]")
    plt.ylabel(ylabel)
    plt.title(title)
    plt.grid(True)
    plt.yscale("log")
    plt.tight_layout()
    plt.savefig(save_path, dpi=300)
    plt.close()

    print(f"Graph saved: {save_path}")


# =========================================================
# Make H average graph
# =========================================================

make_average_graph(
    data=df_avg,
    x_col="time_s_mean",
    y_col="H",
    title="Leak Rate vs Time 0.2 Hz 30-point Average",
    ylabel="Leak Rate [Pa m³/s]",
    save_path=h_graph_path
)


# =========================================================
# Make I average graph
# =========================================================

make_average_graph(
    data=df_avg,
    x_col="time_s_mean",
    y_col="I",
    title="Pressure vs Time 0.2 Hz 30-point Average",
    ylabel="Pressure [Pa]",
    save_path=i_graph_path
)


# =========================================================
# Insert graphs into Excel
# =========================================================

wb = load_workbook(output_excel_path)

if "Graphs" in wb.sheetnames:
    ws_graph = wb["Graphs"]
else:
    ws_graph = wb.create_sheet("Graphs")

ws_graph["A1"] = "Leak Rate vs Time 0.2 Hz 30-point Average"
ws_graph["A25"] = "Pressure vs Time 0.2 Hz 30-point Average"

img_h = Image(str(h_graph_path))
img_i = Image(str(i_graph_path))

img_h.width = 720
img_h.height = 360

img_i.width = 720
img_i.height = 360

ws_graph.add_image(img_h, "A2")
ws_graph.add_image(img_i, "A26")

wb.save(output_excel_path)

print(f"Graphs inserted into Excel: {output_excel_path}")