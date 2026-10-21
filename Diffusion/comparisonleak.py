import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.drawing.image import Image


# =========================================================
# 1. Settings
# =========================================================

# 比較したい2つのCSVファイル
csv1_path = Path("Test202606140001.CSV")
csv2_path = Path("Test202606160001.CSV")

# それぞれの名前。Excelやグラフの凡例に使う
csv1_name = "Test 1"
csv2_name = "Test 2"

# 出力Excelファイル名
output_excel_path = Path("LeakTest_comparison.xlsx")

# グラフ画像名
h_graph_path = Path("comparison_leakrate_H.png")
i_graph_path = Path("comparison_pressure_I.png")

# 平均化秒数
average_sec = 10

# 安定区間だけで統計を取りたい場合に使う
# 例：start_time_s = 3600, end_time_s = 7200
# 全データを使うなら None のままでOK
start_time_s = None
end_time_s = None


# =========================================================
# 2. Excel-like ROUND function
# =========================================================

def excel_round(x, digits=0):
    """
    ExcelのROUNDに近い丸め方。
    Python標準round()は銀行丸めなので、Excelとずれる場合がある。
    """
    factor = 10 ** digits
    return np.sign(x) * np.floor(np.abs(x) * factor + 0.5) / factor


# =========================================================
# 3. Read CSV with encoding fallback
# =========================================================

def read_csv_auto_encoding(path):
    """
    日本語CSV対策。
    shift_jis, cp932, utf-8-sig, utf-8 の順に試す。
    """
    encodings = ["shift_jis", "cp932", "utf-8-sig", "utf-8"]

    last_error = None

    for enc in encodings:
        try:
            df = pd.read_csv(
                path,
                skiprows=14,
                header=None,
                names=["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"],
                encoding=enc
            )
            print(f"{path} read successfully with encoding: {enc}")
            return df

        except UnicodeDecodeError as e:
            last_error = e
            print(f"{path}: failed with encoding {enc}")

    raise last_error


# =========================================================
# 4. Process one CSV
# =========================================================

def process_leak_csv(path, name, average_sec=10):
    """
    CSVを読み込み、F〜K列を計算し、10秒平均データも作る。
    """

    # CSV読み込み
    df = read_csv_auto_encoding(path)

    # 数値化
    for col in ["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    # 読み取り失敗行を削除
    df = df.dropna(
        subset=["time_s", "B_CH1", "C_CH2", "D_CH3", "E_CH4"]
    ).reset_index(drop=True)

    # 実験名を追加
    df["sample"] = name

    # 時間[h]も追加
    df["time_h"] = df["time_s"] / 3600

    # =====================================================
    # Excel F16:K16 の処理
    # =====================================================

    # F = 10 - C - 12
    df["F"] = 10 - df["C_CH2"] - 12

    # G = ROUND(F, 0)
    df["G"] = excel_round(df["F"], 0)

    # H = B × 10^G
    # H：リークレート相当
    df["H"] = df["B_CH1"] * (10 ** df["G"])

    # I = 10^(D - 5.5) × 100
    # I：圧力相当
    df["I"] = (10 ** (df["D_CH3"] - 5.5)) * 100

    # J = 10^(E - 3)
    df["J"] = 10 ** (df["E_CH4"] - 3)

    # K = J - I
    df["K"] = df["J"] - df["I"]

    # =====================================================
    # 10秒平均データ
    # =====================================================

    df_avg = (
        df
        .assign(time_bin=(df["time_s"] // average_sec) * average_sec)
        .groupby("time_bin", as_index=False)
        .agg({
            "B_CH1": "mean",
            "C_CH2": "mean",
            "D_CH3": "mean",
            "E_CH4": "mean",
            "F": "mean",
            "G": "mean",
            "H": "mean",
            "I": "mean",
            "J": "mean",
            "K": "mean"
        })
    )

    df_avg["sample"] = name
    df_avg["time_h"] = df_avg["time_bin"] / 3600

    return df, df_avg


# =========================================================
# 5. Statistics function
# =========================================================

def calc_stats(series):
    """
    バラツキ評価用の統計量を計算する。
    リークレートのような指数スケールの値も考慮して、
    算術平均・中央値・標準偏差・CV・IQR・log10統計を出す。
    """

    s = pd.to_numeric(series, errors="coerce").dropna()
    s_pos = s[s > 0]

    if len(s) == 0:
        return {}

    q1 = s.quantile(0.25)
    q3 = s.quantile(0.75)
    p05 = s.quantile(0.05)
    p95 = s.quantile(0.95)

    mean = s.mean()
    median = s.median()
    std = s.std(ddof=1)
    var = s.var(ddof=1)

    # CV = 標準偏差 / 平均
    if mean != 0:
        cv_percent = std / mean * 100
    else:
        cv_percent = np.nan

    # Median absolute deviation
    mad_median = np.median(np.abs(s - median))

    # RMS
    rms = np.sqrt(np.mean(s ** 2))

    # Modeは連続値ではあまり強くないが、参考として出す
    mode_values = s.mode()
    if len(mode_values) > 0:
        mode = mode_values.iloc[0]
    else:
        mode = np.nan

    # log10統計。HやIが正の値のときだけ計算
    if len(s_pos) > 0:
        log_s = np.log10(s_pos)
        geometric_mean = 10 ** log_s.mean()
        log10_mean = log_s.mean()
        log10_std = log_s.std(ddof=1)
    else:
        geometric_mean = np.nan
        log10_mean = np.nan
        log10_std = np.nan

    stats = {
        "count": len(s),
        "mean": mean,
        "median": median,
        "mode_reference": mode,
        "std": std,
        "variance": var,
        "cv_percent": cv_percent,
        "min": s.min(),
        "max": s.max(),
        "range": s.max() - s.min(),
        "q1_25%": q1,
        "q3_75%": q3,
        "iqr": q3 - q1,
        "p05": p05,
        "p95": p95,
        "p95_minus_p05": p95 - p05,
        "mad_median": mad_median,
        "rms": rms,
        "geometric_mean": geometric_mean,
        "log10_mean": log10_mean,
        "log10_std": log10_std
    }

    return stats


# =========================================================
# 6. Make summary table
# =========================================================

def make_summary(df, sample_name, target_cols):
    """
    H, Iなど指定した列について統計量をまとめる。
    """

    rows = []

    for col in target_cols:
        stats = calc_stats(df[col])

        for stat_name, value in stats.items():
            rows.append({
                "sample": sample_name,
                "parameter": col,
                "stat": stat_name,
                "value": value
            })

    return pd.DataFrame(rows)


# =========================================================
# 7. Compare two summaries
# =========================================================

def compare_summaries(summary1, summary2, name1, name2):
    """
    2つの統計結果を横に並べて、差・比・変化率を計算する。
    """

    s1 = summary1.rename(columns={"value": name1})
    s2 = summary2.rename(columns={"value": name2})

    merged = pd.merge(
        s1[["parameter", "stat", name1]],
        s2[["parameter", "stat", name2]],
        on=["parameter", "stat"],
        how="outer"
    )

    merged["difference_2_minus_1"] = merged[name2] - merged[name1]
    merged["ratio_2_over_1"] = merged[name2] / merged[name1]
    merged["change_percent"] = (merged[name2] - merged[name1]) / merged[name1] * 100

    return merged


# =========================================================
# 8. Optional stable region filter
# =========================================================

def filter_time_range(df, start_time_s=None, end_time_s=None, time_col="time_s"):
    """
    安定区間だけで統計を取りたい場合に使う。
    start_time_s / end_time_s が None なら全データ。
    """

    filtered = df.copy()

    if start_time_s is not None:
        filtered = filtered[filtered[time_col] >= start_time_s]

    if end_time_s is not None:
        filtered = filtered[filtered[time_col] <= end_time_s]

    return filtered


# =========================================================
# 9. Graph function
# =========================================================

def make_comparison_dot_graph(df1, df2, name1, name2, x_col, y_col, title, ylabel, save_path):
    """
    2つのデータを同じグラフにドットで重ねる。
    """

    p1 = df1[[x_col, y_col]].copy().dropna()
    p2 = df2[[x_col, y_col]].copy().dropna()

    # logスケール用に0以下を除外
    p1 = p1[p1[y_col] > 0]
    p2 = p2[p2[y_col] > 0]

    plt.figure(figsize=(10, 5))

    plt.scatter(p1[x_col], p1[y_col], s=8, label=name1)
    plt.scatter(p2[x_col], p2[y_col], s=8, label=name2)

    plt.xlabel("Time [s]")
    plt.ylabel(ylabel)
    plt.title(title)
    plt.grid(True)
    plt.yscale("log")
    plt.legend()

    plt.tight_layout()
    plt.savefig(save_path, dpi=300)
    plt.close()

    print(f"Graph saved: {save_path}")


# =========================================================
# 10. Main process
# =========================================================

# CSV読み込み・F〜K計算・10秒平均
df1, df1_avg = process_leak_csv(csv1_path, csv1_name, average_sec)
df2, df2_avg = process_leak_csv(csv2_path, csv2_name, average_sec)

# 統計に使う範囲
# 基本は10秒平均後のデータを使う
df1_stats_target = filter_time_range(
    df1_avg,
    start_time_s=start_time_s,
    end_time_s=end_time_s,
    time_col="time_bin"
)

df2_stats_target = filter_time_range(
    df2_avg,
    start_time_s=start_time_s,
    end_time_s=end_time_s,
    time_col="time_bin"
)

# HとIについて統計を出す
target_cols = ["H", "I"]

summary1 = make_summary(df1_stats_target, csv1_name, target_cols)
summary2 = make_summary(df2_stats_target, csv2_name, target_cols)

comparison = compare_summaries(summary1, summary2, csv1_name, csv2_name)

# 2つの10秒平均データを縦に結合
avg_combined = pd.concat([df1_avg, df2_avg], ignore_index=True)


# =========================================================
# 11. Save to Excel
# =========================================================

with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
    df1.to_excel(writer, index=False, sheet_name=f"{csv1_name}_raw")
    df2.to_excel(writer, index=False, sheet_name=f"{csv2_name}_raw")

    df1_avg.to_excel(writer, index=False, sheet_name=f"{csv1_name}_10s")
    df2_avg.to_excel(writer, index=False, sheet_name=f"{csv2_name}_10s")
    avg_combined.to_excel(writer, index=False, sheet_name="Combined_10s")

    summary1.to_excel(writer, index=False, sheet_name=f"{csv1_name}_stats")
    summary2.to_excel(writer, index=False, sheet_name=f"{csv2_name}_stats")
    comparison.to_excel(writer, index=False, sheet_name="Comparison")

print(f"Excel saved: {output_excel_path}")


# =========================================================
# 12. Make comparison graphs
# =========================================================

make_comparison_dot_graph(
    df1=df1_avg,
    df2=df2_avg,
    name1=csv1_name,
    name2=csv2_name,
    x_col="time_bin",
    y_col="H",
    title=f"Leak Rate Comparison {average_sec}-s Average",
    ylabel="Leak rate H",
    save_path=h_graph_path
)

make_comparison_dot_graph(
    df1=df1_avg,
    df2=df2_avg,
    name1=csv1_name,
    name2=csv2_name,
    x_col="time_bin",
    y_col="I",
    title=f"Pressure Comparison {average_sec}-s Average",
    ylabel="Pressure I",
    save_path=i_graph_path
)


# =========================================================
# 13. Insert graphs into Excel
# =========================================================

wb = load_workbook(output_excel_path)

if "Graphs" in wb.sheetnames:
    ws_graph = wb["Graphs"]
else:
    ws_graph = wb.create_sheet("Graphs")

ws_graph["A1"] = f"Leak Rate Comparison {average_sec}-s Average"
ws_graph["A25"] = f"Pressure Comparison {average_sec}-s Average"

img_h = Image(str(h_graph_path))
img_i = Image(str(i_graph_path))

img_h.width = 720
img_h.height = 360
img_i.width = 720
img_i.height = 360

ws_graph.add_image(img_h, "A2")
ws_graph.add_image(img_i, "A26")

wb.save(output_excel_path)

print(f"Graphs inserted into Excel: {output_excel_path}")


# =========================================================
# 14. Print important comparison
# =========================================================

print("\n===== Main Comparison =====")
main_stats = comparison[
    comparison["stat"].isin([
        "mean",
        "median",
        "std",
        "cv_percent",
        "iqr",
        "p95_minus_p05",
        "geometric_mean",
        "log10_std"
    ])
]

print(main_stats)